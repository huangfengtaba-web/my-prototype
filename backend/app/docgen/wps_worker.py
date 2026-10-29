# -*- coding: utf-8 -*-
"""WPS COM 单 worker：串行执行，避免 WPS 单实例冲突。

硬要求（docs/09_FUSION_PLAN.md §P3-2）：
1. 单线程 + 队列 + pythoncom.CoInitialize()（COM 非线程安全）
2. 进程内 import cli_anything.wps.core；禁止 subprocess 逐条调 CLI
3. 超时 60s → taskkill /F /IM wps.exe 重试 1 次 → 仍失败按 kind 降级
   （docx→python-docx、pptx→python-pptx，其余置 failed）
4. DOCGEN_ENABLED=0 时 submit 直接抛错
"""
import os
import queue
import subprocess
import threading
import time
import uuid

_Q: "queue.Queue[dict]" = queue.Queue(maxsize=8)
_RESULT: dict = {}
_LOCK = threading.Lock()
_JOB_TIMEOUT_S = 60


def _kill_wps():
    try:
        subprocess.run(["taskkill", "/F", "/IM", "wps.exe"], capture_output=True, timeout=10)
    except Exception:
        pass


def _run_with_timeout(fn, timeout_s):
    """在辅助线程执行 fn，超时抛 TimeoutError（COM 调用无法中断，只能放弃等待）。

    runner 线程必须自己 CoInitialize——worker 线程的 COM 套间不跨线程。
    """
    box = {}

    def runner():
        try:
            import pythoncom
            pythoncom.CoInitialize()
        except Exception:
            pass
        try:
            box["ok"] = fn()
        except Exception as e:                      # noqa: BLE001
            box["err"] = e

    t = threading.Thread(target=runner, daemon=True)
    t.start()
    t.join(timeout_s)
    if t.is_alive():
        raise TimeoutError(f"渲染超时（>{timeout_s}s），疑似 WPS COM 挂起")
    if "err" in box:
        raise box["err"]
    return box.get("ok")


def _fallback(kind, summary, out_dir, filename):
    """WPS 两次失败后的降级产物（headless 库，无条件可用）。"""
    if kind == "docx":
        return _fallback_docx(summary, out_dir, filename)
    if kind == "pptx":
        return _fallback_pptx(summary, out_dir, filename)
    raise RuntimeError("WPS 不可用且该类型无 headless 降级")


def _fallback_docx(summary, out_dir, filename):
    from docx import Document

    doc = Document()
    milp = summary.get("milp_status") or {}
    doc.add_heading("GTX FSO Litho 排程报告（降级渲染）", level=1)
    doc.add_paragraph(f"MILP: {milp}")
    doc.add_paragraph("；".join(f"{k}：{v}" for k, v in
                               (summary.get("story_check") or {}).items()))
    sched = summary.get("schedule") or []
    if sched:
        tbl = doc.add_table(rows=1, cols=4)
        for i, h in enumerate(["LotId", "Eqp", "Start", "End"]):
            tbl.rows[0].cells[i].text = h
        for r in sched:
            cells = tbl.add_row().cells
            cells[0].text = str(r.get("lot"))
            cells[1].text = str(r.get("eqp") or "-")
            cells[2].text = str(r.get("start") or "-")
            cells[3].text = str(r.get("end") or "-")
    path = os.path.abspath(os.path.join(out_dir, f"{filename}.docx"))
    doc.save(path)
    return path


def _fallback_pptx(summary, out_dir, filename):
    from pptx import Presentation

    prs = Presentation()
    milp = summary.get("milp_status") or {}
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "GTX FSO Litho 排程汇报（降级渲染）"
    s1.placeholders[1].text = f"MILP ②{milp.get('step2', '-')} ③{milp.get('step3', '-')}"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Scanner Schedule"
    body = s2.placeholders[1].text_frame
    body.text = "Lot / Eqp / 时间"
    for r in (summary.get("schedule") or [])[:8]:
        p = body.add_paragraph()
        p.text = f"{r.get('lot')}  {r.get('eqp') or '-'}  {r.get('start')}~{r.get('end')}"
    path = os.path.abspath(os.path.join(out_dir, f"{filename}.pptx"))
    prs.save(path)
    return path


def _loop():
    try:
        import pythoncom
        pythoncom.CoInitialize()
    except Exception as e:
        print(f"[wps-worker] CoInitialize 失败: {e}")

    while True:
        job = _Q.get()
        if job is None:
            break
        jid = job["job_id"]
        kind, args = job.get("kind"), job.get("args")
        fn = job["fn"]
        try:
            try:
                path = _run_with_timeout(fn, _JOB_TIMEOUT_S)
            except Exception:
                _kill_wps()                             # 第 1 次失败：杀 WPS 重试
                time.sleep(2)
                try:
                    path = _run_with_timeout(fn, _JOB_TIMEOUT_S)
                except Exception as e2:
                    # 两次失败：headless 降级（python-docx / python-pptx）
                    print(f"[wps-worker] WPS 两次失败，降级渲染: {e2!r}")
                    try:
                        path = _fallback(kind, args[0], args[1], args[2])
                    except Exception as e3:
                        import traceback
                        raise RuntimeError(
                            f"降级渲染也失败: {e3!r}\n{traceback.format_exc()[-800:]}") from e3
            with _LOCK:
                _RESULT[jid] = {"status": "done", "path": path, "err": None,
                                "finished_at": time.time()}
        except Exception as e:
            with _LOCK:
                _RESULT[jid] = {"status": "failed", "path": None,
                                "err": repr(e), "finished_at": time.time()}
            _kill_wps()
        finally:
            _Q.task_done()


threading.Thread(target=_loop, daemon=True, name="wps-worker").start()


def submit(fn, timeout=60, kind=None, args=None) -> str:
    if os.getenv("DOCGEN_ENABLED", "1") != "1":
        raise RuntimeError("DOCGEN_DISABLED: DOCGEN_ENABLED=0")
    jid = uuid.uuid4().hex[:12]
    with _LOCK:
        _RESULT[jid] = {"status": "pending", "path": None, "err": None, "finished_at": None}
    _Q.put({"job_id": jid, "fn": fn, "kind": kind, "args": args})
    return jid


def status(jid: str):
    with _LOCK:
        return _RESULT.get(jid)
